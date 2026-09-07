from pptx import Presentation

pptx_path = r'F:\PythonProjects\aml_rocm_native\docs\pre_defense\predefense.pptx'
prs = Presentation(pptx_path)

found = False

def replace_in_shape(shape):
    global found
    if shape.has_text_frame and "master" in shape.text_frame.text.lower():
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if "master" in r.text.lower():
                    r.text = r.text.replace("master's", "B.Sc.").replace("Master's", "B.Sc.").replace("master", "B.Sc.").replace("Master", "B.Sc.")
                    found = True
    elif shape.shape_type == 6:  # group
        for sub in shape.shapes:
            replace_in_shape(sub)

for idx, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        replace_in_shape(shape)

if found:
    prs.save(pptx_path)
    print("Successfully updated predefense.pptx from master's to B.Sc.!")
else:
    print("Could not find master in any run.")
