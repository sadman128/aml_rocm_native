import sys
from pptx import Presentation

pptx_path = r'F:\PythonProjects\aml_rocm_native\docs\pre_defense\predefense.pptx'
prs = Presentation(pptx_path)

print(f"Total slides: {len(prs.slides)}")
print(f"Slide dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches\n")

for idx, slide in enumerate(prs.slides):
    print(f"\n{'='*30} SLIDE {idx + 1} {'='*30}")
    if slide.shapes.title and slide.shapes.title.text:
        print(f"Title: {slide.shapes.title.text.strip()}")
    else:
        print("Title: (None or custom shape)")

    for s_idx, shape in enumerate(slide.shapes):
        # check text
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                print(f"  [Text {s_idx} | {shape.name}]:\n    {txt.replace(chr(10), chr(10) + '    ')}")
        elif shape.has_table:
            print(f"  [Table {s_idx} | {shape.name}]:")
            for r_idx, row in enumerate(shape.table.rows):
                vals = [c.text.strip().replace("\n", " ") for c in row.cells]
                print(f"    R{r_idx}: {' | '.join(vals)}")
        elif shape.shape_type == 6:  # group
            print(f"  [Group {s_idx} | {shape.name}]:")
            for sub_idx, sub in enumerate(shape.shapes):
                if sub.has_text_frame and sub.text_frame.text.strip():
                    sub_txt = sub.text_frame.text.strip().replace(chr(10), " ")
                    print(f"      sub_{sub_idx}: {sub_txt}")
