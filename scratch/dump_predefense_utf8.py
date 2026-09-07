import sys
from pptx import Presentation

pptx_path = r'F:\PythonProjects\aml_rocm_native\docs\pre_defense\predefense.pptx'
prs = Presentation(pptx_path)

out_file = r'F:\PythonProjects\aml_rocm_native\scratch\predefense_dump.txt'

with open(out_file, 'w', encoding='utf-8') as f:
    f.write(f"Total slides: {len(prs.slides)}\n")
    f.write(f"Slide dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches\n\n")

    for idx, slide in enumerate(prs.slides):
        f.write(f"\n{'='*30} SLIDE {idx + 1} {'='*30}\n")
        if slide.shapes.title and slide.shapes.title.text:
            f.write(f"Title: {slide.shapes.title.text.strip()}\n")
        else:
            f.write("Title: (None or custom shape)\n")

        for s_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    f.write(f"  [Text {s_idx} | {shape.name}]:\n    {txt.replace(chr(10), chr(10) + '    ')}\n")
            elif shape.has_table:
                f.write(f"  [Table {s_idx} | {shape.name}]:\n")
                for r_idx, row in enumerate(shape.table.rows):
                    vals = [c.text.strip().replace("\n", " ") for c in row.cells]
                    f.write(f"    R{r_idx}: {' | '.join(vals)}\n")
            elif shape.shape_type == 6:  # group
                f.write(f"  [Group {s_idx} | {shape.name}]:\n")
                for sub_idx, sub in enumerate(shape.shapes):
                    if hasattr(sub, 'text_frame') and sub.text_frame and sub.text_frame.text.strip():
                        sub_txt = sub.text_frame.text.strip().replace(chr(10), " ")
                        f.write(f"      sub_{sub_idx}: {sub_txt}\n")
                    elif hasattr(sub, 'has_text_frame') and sub.has_text_frame and sub.text_frame.text.strip():
                        sub_txt = sub.text_frame.text.strip().replace(chr(10), " ")
                        f.write(f"      sub_{sub_idx}: {sub_txt}\n")
                    # If subgroup
                    elif getattr(sub, 'shape_type', None) == 6:
                        for ssub_idx, ssub in enumerate(sub.shapes):
                            if hasattr(ssub, 'has_text_frame') and ssub.has_text_frame and ssub.text_frame.text.strip():
                                ssub_txt = ssub.text_frame.text.strip().replace(chr(10), " ")
                                f.write(f"        sub_{sub_idx}_{ssub_idx}: {ssub_txt}\n")

print("Dump completed successfully!")
