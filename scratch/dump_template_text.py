import os
from pptx import Presentation

template_dir = r"F:\PythonProjects\aml_rocm_native\docs\pre_defense\slide templates"
for fname in sorted(os.listdir(template_dir)):
    if not fname.endswith(".pptx"): continue
    print(f"\n{'='*60}\n{fname}\n{'='*60}")
    prs = Presentation(os.path.join(template_dir, fname))
    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} ---")
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.replace('\n', ' ').strip()
                if text: texts.append(text)
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text = cell.text.replace('\n', ' ').strip()
                        if text: texts.append(text)
            if shape.shape_type == 6: # Group
                for shp in shape.shapes:
                    if shp.has_text_frame:
                        text = shp.text.replace('\n', ' ').strip()
                        if text: texts.append(text)
        print(" | ".join(texts[:20]))
